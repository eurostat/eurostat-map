"""Extract URE 2026 demographic-development maps from workbook sheets M02-M05."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook

URE_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = (
    URE_ROOT
    / "2026/data/Demographic_developments_in_rural_regions_and_areas_URE2026 - for GISCO sent 20260903.xlsx"
)

MAPS = {
    "M02": "population-change.json",
    "M03": "old-age-dependency-ratio.json",
    "M04": "natural-population-change.json",
    "M05": "net-migration.json",
}

LEGEND_LABELS = {
    "M02": "Population change",
    "M03": "Old-age dependency ratio",
    "M04": "Natural change",
    "M05": "Net migration",
}

# The workbook's 3 x 3 legend uses red for predominantly urban, ochre for
# intermediate, and green for predominantly rural regions. Within each hue,
# dark/medium/light represent the three ordered value classes.
COLORS_BY_CLASS = {
    "1": "#4d1318",
    "2": "#e04040",
    "3": "#ffa3a3",
    "4": "#4d3c03",
    "5": "#b09120",
    "6": "#efd18c",
    "7": "#1a501a",
    "8": "#33a033",
    "9": "#a4e1a4",
    ":": "#adadad",
}


def clean_text(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else None


def json_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.strip() == ":":
        return ":"
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value).strip()


def next_i_value(ws, h_label):
    for row in range(1, ws.max_row + 1):
        if clean_text(ws.cell(row, 8).value) == h_label:
            for candidate in range(row, min(row + 7, ws.max_row + 1)):
                value = clean_text(ws.cell(candidate, 9).value)
                if value:
                    return value
    raise ValueError(f"Could not find metadata following H={h_label!r} in {ws.title}")


def thresholds_from_labels(ws):
    labels = [clean_text(ws.cell(row, 12).value) for row in range(15, 18)]
    thresholds = []
    for label in (labels[0], labels[2]):
        matches = re.findall(r"-?\d+(?:\.\d+)?", label or "")
        if not matches:
            raise ValueError(f"Could not extract a threshold from {label!r} in {ws.title}")
        thresholds.append(float(matches[-1]))
    return sorted(thresholds)


def extract_sheet(ws):
    values = {}
    urban_rural = {}
    for nuts, _, value, region_type in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        nuts = clean_text(nuts)
        if not nuts:
            break
        value = json_value(value)
        region_type = json_value(region_type)
        if value is not None:
            values[nuts] = value
        if region_type is not None:
            urban_rural[nuts] = region_type

    # M15:O17 is ordered value-class high-to-low by region type urban-to-rural.
    # The renderer's matrix is indexed region type first and ascending value class second.
    workbook_matrix = [[int(ws.cell(row, col).value) for col in range(13, 16)] for row in range(15, 18)]
    class_matrix = [[workbook_matrix[value_class][region_type] for value_class in (2, 1, 0)] for region_type in range(3)]

    title = clean_text(ws["I6"].value)
    if title and title.startswith(f"Map {int(ws.title[1:])}:"):
        title = title.split(":", 1)[1].strip()
    title = title.replace(", by urban-rural typology", "")

    value_class_labels = [clean_text(ws.cell(row, 12).value) for row in range(17, 14, -1)]
    value_class_labels = [re.split(r"\s+\(", label, maxsplit=1)[0] for label in value_class_labels]

    return {
        "sheet": ws.title,
        "nutsLevel": "mixed",
        "nutsYear": 2024,
        "scale": "60M",
        "chapterTitle": clean_text(ws["I4"].value),
        "title": title,
        "subtitle": clean_text(ws["I7"].value),
        "valueLegendLabel": LEGEND_LABELS[ws.title],
        "unitText": "‰" if "‰" in (clean_text(ws["I10"].value) or "") else "%",
        "urbanRuralLegendLabel": clean_text(ws["M12"].value),
        "thresholds": thresholds_from_labels(ws),
        "valueClassLabels": value_class_labels,
        "urbanRuralLabels": {
            str(region_type): clean_text(ws.cell(14, 12 + region_type).value)
            for region_type in range(1, 4)
        },
        "colorsByClass": COLORS_BY_CLASS,
        "classMatrix": class_matrix,
        "footnote": next_i_value(ws, "Footnotes:"),
        "source": next_i_value(ws, "Sources:"),
        "stats": {"value": values, "urbanRuralType": urban_rural},
    }


def main():
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    for sheet_name, filename in MAPS.items():
        result = extract_sheet(workbook[sheet_name])
        output = URE_ROOT / "2026/data" / filename
        with output.open("w", encoding="utf-8") as output_file:
            json.dump(result, output_file, ensure_ascii=False, indent=2)
        print(f"Wrote {output.relative_to(URE_ROOT)} from {sheet_name}: {len(result['stats']['value'])} values")


if __name__ == "__main__":
    main()
